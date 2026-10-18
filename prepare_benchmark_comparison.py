#!/usr/bin/env python3
"""Direct-curated three-model PTM comparison: full and balanced-negative Excel sheets.

This script DOES NOT read benchmark_results/metrics.csv or benchmark_sites.csv.
It constructs labels directly from the current curated CSVs, joins the three
continuous prediction files on (uid, 1-based site, PTM_type), then computes
all five metrics from scratch.

The two identically formatted, plain (no colours) sheets are:
    Full negatives     All common scored, amino-acid-eligible sites.
    Balanced negatives All common scored positive sites + N+ presumed negatives,
                       drawn WITH replacement for each of 1000 repetitions.
                       Entries are the arithmetic mean over those repetitions.
All three models use exactly the same examples in each repetition.

Run from the repository root:
    python prepare_benchmark_comparison.py \\
      --prediction NRPTMpred=results_newwwww/correct_predictions.csv \\
      --prediction MIND=results_original/correct_predictions_newlabels.csv \\
      --prediction MusiteDeep=results_musitedeep/correct_predictions.csv \\
      --balanced-repeats 1000 --seed 252

Inputs default to:
    data/curated/NR_Sequences_New.csv
    data/curated/NR_InputData_New.csv

Output defaults to:
    benchmark_results/NR_PTMpred_three_model_comparison.xlsx

Dependencies: numpy and openpyxl; install with:
    python -m pip install numpy openpyxl

CAUTION: this is a comparison on the curated annotation population (including
training annotations), NOT an independent test on previously unseen proteins.
Unannotated eligible positions are PRESUMED negatives, not verified negatives.
"""

import argparse
import csv
import math
import re
from collections import Counter
from pathlib import Path

import numpy as np


PTM_TO_AMINO_ACID = {
    "Acetylation_K": "K",
    "Methylation_K": "K",
    "Phosphorylation_S": "S",
    "Phosphorylation_T": "T",
    "Phosphorylation_Y": "Y",
    "Sumoylation_K": "K",
    "Ubiquitination_K": "K",
}
MODELS = ("NRPTMpred", "MIND", "MusiteDeep")
DISPLAY = ("NRPTMpred", "MIND-S", "MusiteDeep")
METRICS = ("Accuracy", "Precision", "Recall", "PR-AUC", "ROC-AUC")
KEY_COLUMNS = ("uid", "site", "PTM_type")


def csv_rows(path, required):
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Input file not found: {path}")
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        missing = set(required) - set(reader.fieldnames or [])
        if missing:
            raise ValueError(f"{path}: missing columns {sorted(missing)}")
        yield from reader


def load_sequences(path):
    sequences = {}
    for row in csv_rows(path, ("UniprotID", "Sequence")):
        uid = row["UniprotID"].strip()
        sequence = row["Sequence"].strip().upper()
        if not uid or not sequence:
            raise ValueError(f"{path}: empty sequence or accession")
        if uid in sequences:
            raise ValueError(f"{path}: duplicate accession {uid}")
        sequences[uid] = sequence
    if not sequences:
        raise ValueError(f"{path}: no sequences found")
    return sequences


def load_positives(path, sequences):
    positives = set()
    excluded = Counter()
    for row in csv_rows(path, ("UniprotID", "Site", "PTM_Type")):
        uid = row["UniprotID"].strip()
        raw_site = row["Site"].strip()
        m = re.fullmatch(r"([A-Za-z])(\d+)", raw_site)
        if not m:
            raise ValueError(f"{path}: invalid site {raw_site!r} for {uid}")
        aa, position = m.group(1).upper(), int(m.group(2))
        ptm = f"{row['PTM_Type'].strip()}_{aa}"
        if ptm not in PTM_TO_AMINO_ACID:
            excluded[ptm] += 1
            continue
        sequence = sequences.get(uid)
        if sequence is None:
            raise ValueError(f"{path}: positive {uid} absent from sequences table")
        if not 1 <= position <= len(sequence) or sequence[position - 1] != aa:
            raise ValueError(
                f"{path}: residue mismatch at {uid}:{raw_site} ({ptm}); "
                "fix curated inputs before evaluating"
            )
        key = (uid, position, ptm)
        if key in positives:
            raise ValueError(f"{path}: duplicate curated positive {key}")
        positives.add(key)
    counts = Counter(p for _, _, p in positives)
    if not positives or any(counts.get(p, 0) == 0 for p in PTM_TO_AMINO_ACID):
        raise ValueError("Curated annotations must contain all seven PTM classes")
    print(f"Curated positives: {len(positives)} across seven PTM classes; "
          f"{sum(excluded.values())} out-of-scope annotations ignored")
    for ptm in PTM_TO_AMINO_ACID:
        print(f"  {ptm:20} N+ = {counts[ptm]}")
    return positives


def construct_candidates(sequences, positives):
    # Unique PTM-residue keys; the same K can be eligible separately for four PTMs.
    labels = {ptm: {} for ptm in PTM_TO_AMINO_ACID}
    aa_to_ptms = {}
    for ptm, aa in PTM_TO_AMINO_ACID.items():
        aa_to_ptms.setdefault(aa, []).append(ptm)
    for uid, seq in sequences.items():
        for site, aa in enumerate(seq, start=1):
            for ptm in aa_to_ptms.get(aa, ()):
                key = (uid, site, ptm)
                labels[ptm][key] = int(key in positives)
    for key in positives:
        if key not in labels[key[2]]:
            raise ValueError(f"Curated positive absent from eligible residues: {key}")
    print(f"Candidate PTM-residue examples: {sum(len(d) for d in labels.values()):,}")
    return labels


def canonical_model(name):
    aliases = {
        "NRPTMpred": "NRPTMpred", "NRPTMKB": "NRPTMpred",
        "MIND": "MIND", "MIND-S": "MIND",
        "MusiteDeep": "MusiteDeep",
    }
    if name.strip() not in aliases:
        raise ValueError(f"Unknown model name {name!r}")
    return aliases[name.strip()]


def prediction_paths(arguments):
    found = {}
    for item in arguments:
        name, sep, filename = item.partition("=")
        if not sep or not name.strip() or not filename.strip():
            raise ValueError(f"Invalid --prediction {item!r}; expected MODEL=CSV")
        name = canonical_model(name)
        if name in found:
            raise ValueError(f"Duplicate --prediction for {name}")
        found[name] = filename.strip()
    if set(found) != set(MODELS):
        raise ValueError(f"Need prediction CSVs for all three models: {MODELS}; "
                         f"provided {sorted(found)}")
    return found


def load_predictions(model, path, sequences, labels):
    scores = {ptm: {} for ptm in PTM_TO_AMINO_ACID}
    outside = 0
    for row in csv_rows(path, (*KEY_COLUMNS, "pred_score")):
        ptm = row["PTM_type"].strip()
        if ptm not in PTM_TO_AMINO_ACID:
            outside += 1
            continue
        uid = row["uid"].strip()
        if uid not in sequences:
            outside += 1
            continue
        try:
            site = int(row["site"].strip())
        except (ValueError, AttributeError) as exc:
            raise ValueError(f"{model}: invalid position {row['site']!r}") from exc
        key = (uid, site, ptm)
        if key not in labels[ptm]:
            raise ValueError(f"{model}: predicted {key}, which is not an eligible "
                             "residue/PTM position in the provided sequences")
        if key in scores[ptm]:
            raise ValueError(f"{model}: duplicate prediction for {key}")
        try:
            value = float(row["pred_score"])
        except (ValueError, TypeError) as exc:
            raise ValueError(f"{model}: invalid probability for {key}") from exc
        if not math.isfinite(value) or value < 0 or value > 1:
            raise ValueError(f"{model}: probability outside [0, 1] for {key}")
        scores[ptm][key] = value
    total = sum(map(len, scores.values()))
    print(f"{model:12} {total:,} valid candidate scores, "
          f"{outside:,} out-of-scope/external rows ignored")
    if not total:
        raise ValueError(f"{model}: no eligible predictions")
    return scores


def ranking_metrics(y, p):
    """Exact average precision and ROC-AUC, including properly grouped score ties.

    This avoids thousands of slow sklearn calls during balanced resampling.
    Algebraically equivalent to sklearn.average_precision_score and roc_auc_score.
    """
    n = len(y)
    n_pos = int(np.sum(y))
    n_neg = n - n_pos
    if n_pos == 0 or n_neg == 0:
        raise ValueError("PR-AUC and ROC-AUC require both positive and negative sites")
    order = np.argsort(-p, kind="mergesort")
    sorted_scores = p[order]
    tp = np.cumsum(y[order], dtype=np.int64)
    # Evaluate only the last element of each distinct score (ties together).
    ends = np.r_[np.flatnonzero(np.diff(sorted_scores) != 0), n - 1]
    tp_at = tp[ends]
    fp_at = ends + 1 - tp_at
    increases = np.diff(np.r_[0, tp_at])
    precision = tp_at / (ends + 1)
    ap = float(np.dot(increases / n_pos, precision))
    # Trapezoidal ROC integration across tied-score threshold groups.
    tp_before = np.r_[0, tp_at[:-1]]
    fp_increases = np.diff(np.r_[0, fp_at])
    roc = float(np.dot((tp_before + tp_at) / (2 * n_pos),
                       fp_increases / n_neg))
    return ap, roc


def metrics(y, p, threshold):
    predicted_positive = p >= threshold
    actual_positive = y.astype(bool)
    tp = int(np.count_nonzero(predicted_positive & actual_positive))
    fp = int(np.count_nonzero(predicted_positive & ~actual_positive))
    fn = int(np.count_nonzero(~predicted_positive & actual_positive))
    tn = len(y) - tp - fp - fn
    ap, roc = ranking_metrics(y, p)
    return np.asarray([
        (tp + tn) / len(y),
        tp / (tp + fp) if tp + fp else 0.0,
        tp / (tp + fn) if tp + fn else 0.0,
        ap, roc,
    ], dtype=np.float64)


def evaluate_all(labels, scores_by_model, threshold, repeats, seed):
    rng = np.random.default_rng(seed)
    full = {}
    balanced = {}
    for ptm in PTM_TO_AMINO_ACID:
        truth = labels[ptm]
        # Match every model on exactly the same (uid,site,PTM_class) examples.
        common = set(truth)
        for model in MODELS:
            common.intersection_update(scores_by_model[model][ptm])
        if not common:
            raise ValueError(f"{ptm}: no scored sites common to all three models")
        keys = sorted(common)
        y = np.fromiter((truth[k] for k in keys), count=len(keys), dtype=np.int8)
        positive_idx = np.flatnonzero(y == 1)
        negative_idx = np.flatnonzero(y == 0)
        n_pos, n_neg = len(positive_idx), len(negative_idx)
        if n_pos == 0 or n_neg == 0:
            raise ValueError(f"{ptm}: common scored sites have only one class")
        available_positives = sum(truth.values())
        missing = len(truth) - len(keys)
        print(f"{ptm:20} common={len(keys):5}, N+={n_pos:4}, "
              f"N-={n_neg:5}, missing from >=1 model={missing:4} "
              f"(including {available_positives - n_pos} positives)")
        probabilities = {
            model: np.fromiter((scores_by_model[model][ptm][k] for k in keys),
                               count=len(keys), dtype=np.float64)
            for model in MODELS
        }
        for model in MODELS:
            full[(ptm, model)] = metrics(y, probabilities[model], threshold)
        sums = {model: np.zeros(len(METRICS), dtype=np.float64) for model in MODELS}
        # True negative bootstrap: every realisation draws N+ negatives WITH
        # replacement. Positives are held fixed; all models share the draw.
        for _ in range(repeats):
            sampled_neg = rng.choice(negative_idx, size=n_pos, replace=True)
            indices = np.concatenate((positive_idx, sampled_neg))
            yy = y[indices]
            for model in MODELS:
                sums[model] += metrics(yy, probabilities[model][indices], threshold)
        for model in MODELS:
            balanced[(ptm, model)] = sums[model] / repeats
    return full, balanced


def write_excel(path, full, balanced):
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise RuntimeError("Install openpyxl: python -m pip install openpyxl") from exc
    wb = Workbook()
    ws = wb.active
    ws.title = "Full negatives"
    wb.create_sheet("Balanced negatives")
    for sheet, results in ((wb.worksheets[0], full),
                           (wb.worksheets[1], balanced)):
        sheet.merge_cells("A1:A2")
        sheet["A1"] = "PTM"
        sheet.column_dimensions["A"].width = 24
        for metric_i, metric_name in enumerate(METRICS):
            first = 2 + 3 * metric_i
            sheet.merge_cells(start_row=1, start_column=first,
                              end_row=1, end_column=first + 2)
            sheet.cell(1, first, metric_name)
            for model_i, display_name in enumerate(DISPLAY):
                column = first + model_i
                sheet.cell(2, column, display_name)
                sheet.column_dimensions[sheet.cell(2, column).column_letter].width = 15
        for row_i, ptm in enumerate(PTM_TO_AMINO_ACID, start=3):
            sheet.cell(row_i, 1, ptm)
            for metric_i in range(len(METRICS)):
                for model_i, model in enumerate(MODELS):
                    cell = sheet.cell(row_i, 2 + metric_i * 3 + model_i,
                                      float(results[(ptm, model)][metric_i]))
                    cell.number_format = "0.000"
        # No fill, colored font, borders, conditional formatting, or themes.
    destination = Path(path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destination)
    return destination


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--sequences", default="data/curated/NR_Sequences_New.csv")
    p.add_argument("--curated", default="data/curated/NR_InputData_New.csv")
    p.add_argument("--prediction", action="append", required=True,
                   metavar="MODEL=CSV", help="Repeat for NRPTMpred, MIND, MusiteDeep")
    p.add_argument("--balanced-repeats", type=int, default=1000)
    p.add_argument("--seed", type=int, default=252)
    p.add_argument("--threshold", type=float, default=0.5)
    p.add_argument("--output", default=(
        "benchmark_results/NR_PTMpred_three_model_comparison.xlsx"))
    args = p.parse_args()
    if args.balanced_repeats < 1:
        p.error("--balanced-repeats must be at least 1")
    if not 0 <= args.threshold <= 1:
        p.error("--threshold must be between 0 and 1")
    paths = prediction_paths(args.prediction)
    sequences = load_sequences(args.sequences)
    positives = load_positives(args.curated, sequences)
    labels = construct_candidates(sequences, positives)
    scores = {name: load_predictions(name, paths[name], sequences, labels)
              for name in MODELS}
    full, balanced = evaluate_all(labels, scores, args.threshold,
                                  args.balanced_repeats, args.seed)
    result = write_excel(args.output, full, balanced)
    print(f"\nCreated: {result}")
    print("Worksheets: Full negatives | Balanced negatives")
    print(f"Balanced: all scored positives + equal number of presumed negatives, "
          f"sampled WITH replacement {args.balanced_repeats} times; fixed seed {args.seed}.")
    print("Both sheets use identical per-PTM common-scored-site populations. "
          "The balanced sheet reports mean values.")
    print("No stored benchmark metrics/site tables were used.")


if __name__ == "__main__":
    main()
